"""입시결과 표 → outcome(11) + evidence(02).

    python3 parse_results.py [--univ id,id] [--doc doc:…] [--force]

원칙 (통합 프롬프트 §5·§6)
- 한 행 = 표의 숫자(또는 기호) 칸 하나. 열 제목 경로(header_path)·행 이름(row_label)·쪽을 근거로 남긴다.
- 집단·측정항목·통계는 열 제목에 **분명한 단어**가 있을 때만 채운다. 애매하면 NULL + classification_status=pending.
  · '최종등록'은 집단 이름이다. 최저라는 뜻이 아니다.
  · 평균/50·70·80%/최저/최고/표준편차는 서로 바꾸지 않는다.
  · 경쟁률·충원인원·충원율·최종예비순위는 별개 측정항목.
- '▨'·'-'·빈칸을 0으로 바꾸지 않는다(빈칸은 행을 만들지 않는다). 원표기는 value_raw.
- 결과 학년도: 문서 본문에서 읽은 연도가 하나면 그 해, 여러 해면 그 쪽/시트에서 한 해만 보일 때만. 아니면 NULL(pending).
- 이미지 표(OCR 텍스트만 있는 쪽)는 표 구조가 없어 여기서 만들지 않는다 → processing_log 에 held.
"""

import argparse
import json
import re
from collections import Counter

from common import RAW, WORK, PARSER_VERSION, connect, text_db, log, h, jdump, detect_years, nfc

EXTRACTOR = "parse_results " + PARSER_VERSION
NUM_RE = re.compile(r"^[\s(]*[-+]?\d[\d,]*(\.\d+)?\s*(%|:1|명|점|등급|대\s*1)?[\s)]*$")
SYMBOLS = {"▨", "-", "―", "–", "*", "※", "x", "X", "비공개", "미공개", "해당없음"}


def cell_str(v):
    if v is None:
        return ""
    return nfc(str(v)).replace("\n", " ").strip()


def is_value(s):
    s = s.strip()
    return bool(s) and (NUM_RE.match(s) is not None or s in SYMBOLS or re.match(r"^\d[\d,.]*\s*:\s*1$", s))


def to_num(s):
    m = re.search(r"[-+]?\d[\d,]*(\.\d+)?", s)
    if not m or s.strip() in SYMBOLS:
        return None
    try:
        return float(m.group(0).replace(",", ""))
    except ValueError:
        return None


# ── 열 제목 → 분류 ──
POP = [(r"1\s*단계", "1단계합격자"), (r"최초\s*합", "최초합격자"), (r"최종\s*등록|등록자", "최종등록자"),
       (r"최종\s*합격", "최종합격자"), (r"지원자", "지원자"), (r"합격자", "합격자(단계 미상)")]
STAT = [(r"(?<![\d.])100\s*%", ("분위값", 100)), (r"(?<![\d.])50\s*%|50%\s*컷|중앙값|\b50\b컷", ("분위값", 50)), (r"(?<![\d.])70\s*%|70%\s*컷", ("분위값", 70)),
        (r"(?<![\d.])80\s*%|80%\s*컷", ("분위값", 80)), (r"(?<![\d.])85\s*%", ("분위값", 85)),
        (r"(?<![\d.])90\s*%", ("분위값", 90)), (r"평균|\bAvg\b|\bAVG\b|\bMean\b", ("평균", None)),
        (r"최저(?!\s*(학력\s*)?기준)(?!\s*충족)(?!\s*통과)|최소|\bLow\b|\bMin\b|컷\s*$", ("최저", None)), (r"최고|최대|\bHigh\b|\bMax\b", ("최고", None)),
        (r"표준\s*편차|\bStd\b|\bSD\b", ("표준편차", None))]
COUNT_METRICS = [(r"경쟁\s*률", "경쟁률"), (r"충원\s*[률율]", "충원율"),
                 (r"최저\s*충족\s*[률율]|충족률|최저\s*통과\s*[률율]", "최저충족률"),
                 (r"최저\s*(학력\s*)?기준\s*통과|최저\s*충족\s*(인원|자)|최저\s*통과\s*(인원|자)", "최저충족인원"),
                 (r"지원\s*[률율]", "지원율"), (r"등록\s*[률율]", "등록률"),
                 (r"최종\s*(예비\s*)?순위|예비\s*(합격\s*)?번호|예비\s*(합격\s*)?순위|추합\s*번호|후보\s*순위|추가\s*통보\s*번호", "최종예비순위"),
                 (r"모집\s*인원|모집$|정원", "모집인원"),
                 (r"지원\s*인원|지원자\s*수|^지원$", "지원인원"), (r"충원\s*(합격)?\s*(인원)?|추가\s*합격|\d+\s*차", "충원인원"),
                 (r"합격\s*인원|합격자\s*수", "합격인원"), (r"등록\s*인원|등록자\s*수", "등록인원")]
SCORE_METRICS = [(r"백분위", "백분위"), (r"표준\s*점수|표점", "표준점수"), (r"환산\s*등급", "환산등급"), (r"등급", "등급"),
                 (r"교과\s*(환산)?|환산\s*점수|환산점|점수|총점", "환산점수"), (r"내신", "내신등급")]


# 열 제목의 마지막 단이 딱 이 말뿐일 때만 = 그 항목의 인원·비율 (전체 제목에 섞여 있으면 쓰지 않는다)
BARE = [(r"지원(자)?(\s*수)?", "지원인원", "지원자"), (r"모집(\s*인원)?|정원", "모집인원", None),
        (r"등록(자)?(\s*수)?|등록\s*기준", "등록인원", "최종등록자"), (r"충원|추가\s*합격|추합", "충원인원", None),
        (r"합격(자)?(\s*수)?", "합격인원", "합격자(단계 미상)"), (r"경쟁(\s*률|율)?", "경쟁률", None),
        (r"입학(자)?(\s*수)?", "등록인원", "최종등록자")]


POP_ONLY = re.compile(r"^(최초\s*합(격)?(자)?|최종\s*등록(자)?|최종\s*합격(자)?|1\s*단계\s*합격(자)?)\s*(\(?\s*(수|인원|명)\s*\)?)?$")
CALC_AVG = re.compile(r"(영역|과목|교과|개)\s*평균|평균\s*등급\s*산출")


# '50% | 70% | 50% | 70%' 처럼 통계 이름만 늘어선 줄 = 헤더의 셋째 단(값이 아니다)
STAT_LABEL = re.compile(r"^\s*(\d{1,3}\s*%\s*(컷|cut|CUT)?|상위\s*\d{1,3}\s*%|평균|최고|최저|최대|최소|중앙값|표준편차|누적|등급|점수|인원|명|배수|Avg|AVG|Max|Min|Std|SD)\s*$", re.I)


def stat_label_row(cells, seen):
    """헤더 아래에 통계 이름만 있는 줄인가.

    HTML rowspan 은 위 칸 글자를 아래 줄에 되풀이해 적는다. 그래서 그 열에서 이미 본 글자는 빼고,
    남은 칸이 전부 통계 이름이면 이 줄도 헤더의 한 단이다. 숫자처럼 보이는 칸('50%')이 있어야 의미가 있다.
    """
    new = [c for i, c in enumerate(cells) if c and c not in seen.get(i, ())]
    if len(new) < 2 or not any(is_value(c) for c in new):
        return False
    return all(STAT_LABEL.match(c) for c in new)


LABEL_HEADER = re.compile(r"학년도|입학\s*년도|학제|주야|수업\s*연한|구분$|계열$"
                          r"|^연번$|^순번$|^번호$|^No\.?$|^NO\.?$|^#$")   # 값이 아니라 행 이름인 열


def classify(header, table_title=None):
    # 점수 계산 방식 설명 속 '평균'(예: 상위 2개영역평균)은 통계가 아니다 → 지우고 판정
    hp = CALC_AVG.sub(" ", header)
    hp = hp + " " + re.sub(r"\s+", "", hp)   # '모 집 인 원'처럼 글자마다 띄어 쓴 제목도 잡는다
    pop = next((v for rx, v in POP if re.search(rx, hp)), None)
    stat = next((v for rx, v in STAT if re.search(rx, hp)), None)
    score = next((v for rx, v in SCORE_METRICS if re.search(rx, hp)), None)
    count = next((v for rx, v in COUNT_METRICS if re.search(rx, hp)), None)
    metric, st, pr = None, None, None
    # 마지막 단이 '지원'·'모집'처럼 한 말뿐이면 그 항목의 인원으로 본다
    last = re.sub(r"\s+", "", CALC_AVG.sub(" ", header.split(" > ")[-1]))
    bare = next(((m, bp) for rx, m, bp in BARE if re.fullmatch(rx, last)), None)
    if bare and not score and not count and not stat:
        count = bare[0]
        pop = pop or bare[1]
    if score and stat:
        metric, st, pr = score, stat[0], stat[1]
    elif count and not stat:
        metric, st = count, "값"
    elif score and not stat:
        metric = score
    elif count:
        metric = count
    elif pop and not stat and POP_ONLY.match(re.sub(r"\s+", " ", header.split(" > ")[-1]).strip()):
        metric, st = "인원", "값"          # 마지막 열 제목이 '최초합'·'최종등록' 처럼 집단 이름뿐일 때만 = 그 집단의 인원
    # 등급·점수 세부
    if metric == "등급":
        metric = "내신등급" if not re.search(r"수능|국어|수학|영어|탐구", hp) else "수능등급"
    if metric == "환산점수" and re.search(r"교과", hp):
        metric = "교과환산점수"
    elif metric == "환산점수" and re.search(r"수능", hp):
        metric = "수능환산점수"
    if metric in ("경쟁률", "충원율", "최저충족률", "최저충족인원", "최종예비순위", "모집인원", "지원인원", "충원인원",
                  "합격인원", "등록인원", "인원"):
        st = "값"
    # 열 제목에 측정항목이 없고 통계만 있으면, 같은 표의 제목에서 측정항목만 찾는다(통계는 표 제목에서 가져오지 않는다)
    if metric is None and stat and table_title and len(table_title) <= 60:
        # 표 제목이 쪽 전체 글자를 긁어온 것이면(60자 초과) 거기서 측정항목을 추측하지 않는다.
        # 경희대 사례: 쪽 전체가 제목으로 잡혀 '학생부 교과 등급 분포' 가 들어가는 바람에
        # '합격자 평균성적'(90.8) 열이 내신등급으로 잘못 분류됐다.
        tt = CALC_AVG.sub(" ", table_title)
        tt = tt + " " + re.sub(r"\s+", "", tt)
        metric = next((v for rx, v in SCORE_METRICS if re.search(rx, tt)), None)
        if metric == "등급":
            metric = "내신등급" if not re.search(r"수능|국어|수학|영어|탐구", tt) else "수능등급"
        elif metric == "환산점수" and re.search(r"교과", tt):
            metric = "교과환산점수"
        if metric:
            st, pr = stat[0], stat[1]
    status = "auto_classified" if metric and st else "pending"
    return pop, metric, st, pr, status


# ── 표 → (헤더 경로, 데이터 행) ──
def split_headers(rows):
    """맨 위 행들 중 값 칸이 거의 없는 행을 헤더로. 여러 열에 걸친 칸은 오른쪽 None 에 퍼뜨린다."""
    hdr, seen = [], {}
    for r in rows:
        cells = [cell_str(c) for c in r]
        vals = sum(1 for c in cells if is_value(c))
        if vals >= max(2, len(cells) * 0.3):
            # 헤더 바로 아래에 '50% | 70%' 같은 통계 이름 줄이 오면 그것도 헤더다
            if not (hdr and stat_label_row(cells, seen)):
                break
        hdr.append(list(r))
        for i, c in enumerate(cells):
            if c:
                seen.setdefault(i, set()).add(c)
        if len(hdr) >= 5:
            break
    ncol = max(len(r) for r in rows) if rows else 0
    paths = [[] for _ in range(ncol)]
    for r in hdr:
        r = list(r) + [None] * (ncol - len(r))
        i = 0
        while i < ncol:
            c = r[i]
            if c is None:
                i += 1
                continue
            j = i + 1
            while j < ncol and r[j] is None:
                j += 1
            span = j - i
            text = cell_str(c)
            parts = distribute(str(c), span)
            for k in range(span):
                t = parts[k] if parts else text
                if t:
                    paths[i + k].append(t)
            i = j
    return len(hdr), [" > ".join(dict.fromkeys(p)) for p in paths]


def distribute(raw, span):
    """'A B A B\\nX Y' 처럼 여러 열 제목이 한 칸에 뭉친 것을 span 개로 나눈다. 못 나누면 None."""
    if span <= 1:
        return None
    lines = [l.split() for l in str(raw).split("\n") if l.strip()]
    if not lines:
        return None
    out = [[] for _ in range(span)]
    ok = False
    for toks in lines:
        if len(toks) % span == 0:
            g = len(toks) // span
            for k in range(span):
                out[k].append(" ".join(toks[k * g:(k + 1) * g]))
            ok = True
        else:
            for k in range(span):
                out[k].append(" ".join(toks))
    return [" ".join(x) for x in out] if ok else None


def lead_value(c):
    """칸 맨 앞 토큰이 값이면 그 토큰('222% 국어국문학과' → '222%'), 아니면 None."""
    t = c.split()
    return t[0] if t and is_value(t[0]) else None


def label_cols(headers, data):
    """값이 아닌 글자가 주로 있는 열 = 행 이름 열."""
    cols = []
    for j, hp in enumerate(headers):
        col = [cell_str(r[j]) if j < len(r) else "" for r in data]
        texts = sum(1 for c in col if c and not lead_value(c))
        vals = sum(1 for c in col if c and lead_value(c))
        if texts > vals:
            cols.append(j)
    return cols


ADM_COL = re.compile(r"전형")
PROG_COL = re.compile(r"모집\s*단위|학과|학부|전공")
GROUP_COL = re.compile(r"^군$|모집\s*군")


def emit_table(con, ctx, headers, data, where):
    """where: dict(page_index / sheet) 근거 위치"""
    lcols = label_cols(headers, data)
    ff = {}
    n = 0
    for ri, r in enumerate(data):
        cells = [cell_str(c) for c in r] + [""] * (len(headers) - len(r))
        # 행 이름(병합된 위 칸은 앞 행 값을 이어받는다)
        labels = {}
        for j in lcols:
            v = cells[j]
            if v and not is_value(v):
                ff[j] = v
                # 오른쪽 열이 새로 바뀌면 그 오른쪽 이어받기를 지운다
                for k in list(ff):
                    if k > j and k in lcols and not cells[k]:
                        ff.pop(k, None)
            labels[j] = ff.get(j, "")
        if not any(labels.values()):
            continue
        last = {j: (headers[j] or "").split(" > ")[-1] for j in lcols}
        adm = next((labels[j] for j in lcols if ADM_COL.search(last[j]) and labels[j]), None) or ctx.get("admission")
        prog = (next((labels[j] for j in lcols if PROG_COL.search(last[j]) and labels[j]), None)
                or next((labels[j] for j in reversed(lcols) if PROG_COL.search(headers[j] or "") and labels[j]), None))
        grp = next((labels[j] for j in lcols if GROUP_COL.search(last[j]) and labels[j]), None)
        row_label = " / ".join(labels[j] for j in lcols if labels[j])
        if re.search(r"^(합계|계|총계|소계)$", (prog or "").strip()):
            agg = "admission"
        else:
            agg = "program" if prog else "admission"
        for j, hp in enumerate(headers):
            if j in lcols or LABEL_HEADER.search(hp or ""):
                continue
            v = cells[j] if j < len(cells) else ""
            if not v:
                continue
            first = v.split()[0]
            if not is_value(first):
                continue
            pop, metric, st, pr, cstat = classify(hp or "", ctx.get("title"))
            hy = detect_years(hp or "")
            cell_year = hy[0] if len(hy) == 1 else (None if len(hy) > 1 else ctx.get("year"))
            fr = re.search(r"(\d+)\s*차", (hp or "").split(" > ")[-1]) if metric == "충원인원" else None
            num = to_num(first)
            sym = first if first.strip() in SYMBOLS else None
            oid = "out:" + h(ctx["doc"], ctx["univ"], where.get("key"), ri, j)
            unit = "%" if first.endswith("%") else (":1" if ":" in first else None)
            con.execute("""INSERT OR REPLACE INTO outcome (outcome_id, result_academic_year, university_id, program_name_raw,
                phase, admission_group, admission_name_raw, aggregation_scope, population_stage, applicant_background, metric,
                component_scope, statistic_type, percentile_rank, value_num, value_raw, unit, suppression_reason, header_path,
                row_label, document_id, classification_status, value_status, review_status, fill_round)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                        (oid, cell_year, ctx["univ"], prog, ctx.get("phase"), grp, adm, agg, pop,
                         ctx.get("background", "구분미공개"), metric,
                         ctx.get("component_scope") if metric in ("내신등급", "교과환산점수", "환산점수", "수능환산점수", "백분위", "표준점수", "수능등급") else None,
                         st, pr,
                         num, v, unit, ctx.get("suppression") if sym == "▨" else None, hp, row_label, ctx["doc"],
                         cstat, "pending" if sym else ("confirmed" if num is not None else "pending"), "auto_validated",
                         fr.group(1) + "차" if fr else None))
            con.execute("""INSERT OR REPLACE INTO evidence (evidence_id, document_id, target_table, target_record_id,
                target_field, page_index, table_title, header_path, row_label, column_label, sheet_cell, raw_text,
                footnote_text, normalized_value, unit, value_status, extraction_method, extractor_version, review_status)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                        ("ev:" + oid[4:], ctx["doc"], "outcome", oid, "value_raw", where.get("page_index"),
                         ctx.get("title"), hp, row_label, (hp or "").split(" > ")[-1], where.get("cell_fn", lambda a, b: None)(ri, j),
                         v, ctx.get("footnotes"), None if num is None else str(num), unit,
                         "pending" if sym else "confirmed", where.get("method"), EXTRACTOR, "auto_validated"))
            n += 1
    return n


def page_context(text):
    lines = [l.strip() for l in (text or "").splitlines() if l.strip()]
    foot = [l for l in lines if re.match(r"^[•※*▶◦·\-]|^\(?주\)?", l) or "▨" in l]
    adm = None
    # 띄어쓰기 없는 한 토막만(앞 문장이 붙지 않게): '학생부교과전형[일반]', '논술(논술우수자)전형' 등
    cands = re.findall(r"(?<![가-힣])([가-힣A-Za-z0-9·()Ⅰ-Ⅻ]{1,24}전형(?:\s?\[[^\]]{1,24}\]|\s?\([^)]{1,24}\))?)", " ".join(lines[:12]))
    cands = [c.strip() for c in cands if 4 <= len(c.strip()) <= 40 and not re.search(r"대학입학전형|입학전형결과|전형결과|나타냄|제외한|포함", c)]
    if len(set(cands)) == 1:
        adm = cands[0]
    comp = None
    ft = " ".join(foot)
    if re.search(r"면접.{0,12}제외", ft):
        comp = "면접 제외 교과성적(원문 각주)"
    supp = None
    m = re.search(r"▨\s*[:：]?\s*([^•※\n]{5,120})", ft)
    if m:
        supp = m.group(1).strip()
    years = detect_years(text)
    return {"footnotes": " | ".join(foot)[:1500] or None, "admission": adm, "component_scope": comp,
            "suppression": supp, "title": " ".join(lines[:3])[:200], "page_years": years}


def year_for(doc_years, page_years):
    if len(doc_years) == 1:
        return doc_years[0]
    if len(page_years) == 1:
        return page_years[0]
    return None


def repaired(path):
    """pdfplumber(pdfminer)가 못 여는 손상 PDF → PyMuPDF 로 복구본을 만들어 그 경로를 돌려준다."""
    import fitz
    out = WORK / "tmp_repair"
    out.mkdir(parents=True, exist_ok=True)
    dst = out / (h(str(path)) + ".pdf")
    if not dst.exists():
        with fitz.open(path) as d:
            d.save(str(dst), clean=True, garbage=3)
    return dst


def parse_pdf(con, tdb, doc, univ, path, phase, doc_years):
    import pdfplumber
    n = 0
    try:
        pdfplumber.open(path).close()
    except Exception:  # noqa: BLE001
        path = repaired(path)
    with pdfplumber.open(path) as pdf:
        for pi, page in enumerate(pdf.pages):
            try:
                tables = page.extract_tables()
            except Exception:  # noqa: BLE001
                continue
            t = tdb.execute("SELECT text FROM page_text WHERE document_id=? AND page_index=?", (doc, pi)).fetchone()
            pc = page_context(t[0] if t else "")
            for ti, tb in enumerate(tables):
                if len(tb) < 3:
                    continue
                nh, headers = split_headers(tb)
                data = tb[nh:]
                if not data or not any(is_value(cell_str(c)) for r in data for c in r):
                    continue
                ctx = dict(pc, doc=doc, univ=univ, phase=phase, year=year_for(doc_years, pc["page_years"]))
                n += emit_table(con, ctx, headers, data, {"key": f"p{pi}t{ti}", "page_index": pi, "method": "table_rule(pdfplumber)"})
    return n


def parse_xlsx(con, doc, univ, path, phase, doc_years):
    import openpyxl
    from openpyxl.utils import get_column_letter
    wb = openpyxl.load_workbook(path, data_only=True)
    n = 0
    for si, ws in enumerate(wb.worksheets):
        grid = [[c.value for c in r] for r in ws.iter_rows()]
        if not grid:
            continue
        # 병합 칸: 왼쪽 위 값만 두고 나머지는 None(헤더 퍼뜨리기와 같은 규칙)
        for mr in ws.merged_cells.ranges:
            for rr in range(mr.min_row, mr.max_row + 1):
                for cc in range(mr.min_col, mr.max_col + 1):
                    if (rr, cc) != (mr.min_row, mr.min_col) and rr - 1 < len(grid) and cc - 1 < len(grid[rr - 1]):
                        v = grid[mr.min_row - 1][mr.min_col - 1]
                        # 가로 병합 = 왼쪽만(헤더 퍼뜨리기), 세로 병합 = 글자만 아래로(숫자는 복제하지 않는다: 정원 공유 등)
                        grid[rr - 1][cc - 1] = (v if (rr != mr.min_row and cc == mr.min_col and not is_value(cell_str(v))) else None)
        # 제목 줄(값 없는 긴 글) 건너뛰기
        start = 0
        while start < len(grid) and sum(1 for c in grid[start] if c not in (None, "")) <= 1:
            start += 1
        rows = grid[start:]
        if len(rows) < 3:
            continue
        pre = " ".join(cell_str(c) for r in grid[:start] for c in r if c)
        pc = page_context(pre + "\n" + ws.title)
        nh, headers = split_headers(rows)
        data = rows[nh:]
        ctx = dict(pc, doc=doc, univ=univ, phase=phase, year=year_for(doc_years, detect_years(pre + ws.title)))
        off = start + nh

        def cell_fn(ri, j, ws=ws, off=off):
            return f"{ws.title}!{get_column_letter(j + 1)}{off + ri + 1}"
        n += emit_table(con, ctx, headers, data, {"key": f"s{si}", "page_index": si, "method": "xlsx_cell", "cell_fn": cell_fn})
    return n


def parse_html(con, doc, univ, path, phase, doc_years):
    from bs4 import BeautifulSoup
    raw = path.read_bytes()
    m = re.search(rb'charset=["\']?([\w-]+)', raw[:3000], re.I)
    try:
        html = raw.decode(m.group(1).decode() if m else "utf-8", errors="ignore")
    except LookupError:
        html = raw.decode("utf-8", errors="ignore")
    soup = BeautifulSoup(html, "html.parser")
    n = 0
    for ti, tb in enumerate(soup.find_all("table")):
        grid = {}
        for ri, tr in enumerate(tb.find_all("tr")):
            ci = 0
            for td in tr.find_all(["td", "th"]):
                while (ri, ci) in grid:
                    ci += 1
                def _span(v):
                    m = re.match(r"\s*(\d+)", str(v or "1"))   # 'rowspan="8 class=..."' 처럼 깨진 값 방어
                    return max(1, min(int(m.group(1)) if m else 1, 200))
                rs, cs = _span(td.get("rowspan", 1)), _span(td.get("colspan", 1))
                txt = td.get_text(" ", strip=True)
                for dr in range(rs):
                    for dc in range(cs):
                        # 가로 병합은 왼쪽만 값(헤더 퍼뜨리기 규칙), 세로 병합은 아래로 복사
                        grid[(ri + dr, ci + dc)] = txt if dc == 0 else None
                ci += cs
        if not grid:
            continue
        nr = max(r for r, _ in grid) + 1
        nc = max(c for _, c in grid) + 1
        rows = [[grid.get((r, c), "") for c in range(nc)] for r in range(nr)]
        if len(rows) < 2:
            continue
        cap = tb.find("caption")
        prev = tb.find_previous(["h1", "h2", "h3", "h4", "strong", "p"])
        title = " ".join(x.get_text(" ", strip=True) for x in (cap, prev) if x)[:300]
        pc = page_context(title)
        nh, headers = split_headers(rows)
        data = rows[nh:]
        ctx = dict(pc, doc=doc, univ=univ, phase=phase, year=year_for(doc_years, detect_years(title)))
        n += emit_table(con, ctx, headers, data, {"key": f"t{ti}", "page_index": 0, "method": "html_table"})
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--univ")
    ap.add_argument("--doc")
    ap.add_argument("--force", action="store_true")
    a = ap.parse_args()
    con, tdb = connect(), text_db()
    q = """SELECT DISTINCT f.document_id, f.university_id, f.raw_path, d.format, d.admission_scope, d.academic_years_detected,
                  d.academic_years_claimed, d.converted_from, d.text_layer
           FROM document_file f JOIN document d USING(document_id) JOIN university_campus u ON u.university_id=f.university_id
           WHERE d.document_type='입시결과' AND u.in_target=1"""
    rows = list(con.execute(q))
    if a.univ:
        rows = [r for r in rows if r["university_id"] in a.univ.split(",")]
    if a.doc:
        rows = [r for r in rows if r["document_id"] == a.doc]
    # 변환본(PDF)의 원본이 HTML·XLS 이면 구조가 살아 있는 원본 쪽을 읽는다
    originals = {r["converted_from"] for r in rows if r["converted_from"]}
    seen = set()
    stats = Counter()
    for r in rows:
        did, uid = r["document_id"], r["university_id"]
        if (did, uid) in seen:
            continue
        seen.add((did, uid))
        if not a.force and con.execute("SELECT 1 FROM processing_log WHERE document_id=? AND stage=? AND parser_version=? AND detail LIKE ?",
                                       (did, "results", PARSER_VERSION, f"{uid}:%")).fetchone():
            continue
        path = RAW / r["raw_path"]
        fmt = r["format"]
        phase = {"수시": "수시", "정시": "정시"}.get(r["admission_scope"])
        dy = json.loads(r["academic_years_detected"] or "[]") or json.loads(r["academic_years_claimed"] or "[]")
        dy_claim = json.loads(r["academic_years_claimed"] or "[]")
        doc_years = dy if dy else dy_claim
        if r["converted_from"]:
            o = con.execute("SELECT format FROM document WHERE document_id=?", (r["converted_from"],)).fetchone()
            if o and o[0] in ("HTML", "XLS", "XLSX"):
                stats["skip_converted(원본으로 처리)"] += 1
                continue
        try:
            con.execute("DELETE FROM evidence WHERE target_table='outcome' AND target_record_id IN (SELECT outcome_id FROM outcome WHERE document_id=? AND university_id=?)", (did, uid))
            con.execute("DELETE FROM outcome WHERE document_id=? AND university_id=?", (did, uid))
            if fmt == "PDF":
                if r["text_layer"] == "image_only":
                    log(con, did, "results", "held", f"{uid}: 이미지 표 — OCR 텍스트만 있어 표 구조 추출 안 함(수작업·렌더 판독 필요)")
                    stats["held_image"] += 1
                    continue
                n = parse_pdf(con, tdb, did, uid, path, phase, doc_years)
            elif fmt == "XLSX":
                n = parse_xlsx(con, did, uid, path, phase, doc_years)
            elif fmt == "HTML":
                n = parse_html(con, did, uid, path, phase, doc_years)
            elif fmt == "XLS":
                stats["skip_xls(변환본 XLSX로)"] += 1
                continue
            else:
                log(con, did, "results", "held", f"{uid}: 형식 {fmt} 표 추출 대상 아님")
                stats["held_format"] += 1
                continue
            log(con, did, "results", "done" if n else "partial", f"{uid}: 값 {n}개" + ("" if n else " — 표를 찾지 못함"))
            stats["done" if n else "no_table"] += 1
            stats["values"] += n
        except Exception as e:  # noqa: BLE001
            log(con, did, "results", "failed", f"{uid}: {type(e).__name__}: {e}"[:500])
            stats["failed"] += 1
        con.commit()
    print(dict(stats))


if __name__ == "__main__":
    main()
