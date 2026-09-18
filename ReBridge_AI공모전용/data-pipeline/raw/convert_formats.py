"""RAW 파일 형식 통일 (2026-09-18 동근님 지시): PDF·XLSX·CSV가 아닌 파일을 PDF나 XLSX로 바꾼다.

- 내용이 바뀌지 않았는지 검사를 통과한 것만 바꾼다. 실패하면 변환본을 버리고 원본을 그대로 둔다.
- 원본은 지우지 않고 같은 학교 폴더의 `_원본형식/` 으로 옮긴다(manifest 경로도 고친다).
- 변환본은 `{원래이름}_변환.pdf|xlsx` 로 두고 manifest `convert.csv` 에 1줄 쓴다.

    XLS·XLSM → XLSX  LibreOffice. 모든 시트·모든 칸 값 비교
    JSON     → XLSX  표로 펼침. 모든 값 비교
    JPG      → PDF   JPEG 바이트를 다시 압축하지 않고 그대로 넣음. pdfimages로 꺼내 바이트 동일 확인
    PNG·WEBP → PDF   무손실(Flate). 픽셀 동일 확인
    HTML     → PDF   Chrome headless 인쇄. HTML에 보이는 숫자·글자가 PDF에 다 있는지 확인
    HWP      → PDF   pyhwp(hwp5html)로 HTML → Chrome 인쇄. HTML 글자 전부 + hwp5txt 본문이 PDF에 있는지 확인
                     (한컴 화면과 줄바꿈·글꼴은 다를 수 있다. 원본 HWP는 _원본형식/ 에 남는다)
    HWPX     → PDF   LibreOffice. 파일 안 XML의 글자가 PDF에 다 있는지 확인
    pyhwp 위치: 환경변수 PYHWP_BIN (기본 ~/.venvs/pyhwp/bin)

    python3 convert_formats.py [--only html,xls,...] [--dry]
"""

import argparse
import csv
import datetime as dt
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zlib
from collections import Counter
from pathlib import Path

import rawlib as R

CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
KEEP = {".pdf", ".xlsx", ".csv"}
KIND = {".xls": "xls", ".xlsm": "xls", ".json": "json", ".jpg": "jpg", ".jpeg": "jpg",
        ".png": "img", ".webp": "img", ".html": "html", ".htm": "html", ".hwp": "hwp", ".hwpx": "hwpx"}
WHO = "convert"
PYHWP = Path(__import__("os").environ.get("PYHWP_BIN", Path.home() / ".venvs/pyhwp/bin"))


# ---------- 변환 ----------
def xls_to_xlsx(src, dst):
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td) / ("in" + src.suffix.lower())
        shutil.copy2(src, tmp)
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", td, str(tmp)],
                       capture_output=True, timeout=180)
        out = Path(td) / "in.xlsx"
        if not out.exists():
            return "LibreOffice 변환 실패"
        shutil.move(out, dst)


def json_to_xlsx(src, dst):
    import pandas as pd
    data = json.loads(src.read_text(encoding="utf-8"))
    rows = data if isinstance(data, list) else next((v for v in data.values() if isinstance(v, list)), [data])
    pd.json_normalize(rows).to_excel(dst, index=False)


def _pdf(width, height, img_obj):
    """이미지 1장을 쪽 1장에 원래 크기(1px=1pt)로 넣은 최소 PDF."""
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {width} {height}] "
        f"/Resources << /XObject << /Im0 4 0 R >> >> /Contents 5 0 R >>".encode(),
        img_obj,
    ]
    content = f"q {width} 0 0 {height} 0 0 cm /Im0 Do Q".encode()
    objs.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offs = []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    x = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offs)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF\n".encode()
    return bytes(out)


def jpg_to_pdf(src, dst):
    from PIL import Image
    im = Image.open(src)
    if im.mode not in ("RGB", "L", "CMYK"):
        return f"JPEG 색 모드 {im.mode}"
    cs = {"RGB": "/DeviceRGB", "L": "/DeviceGray", "CMYK": "/DeviceCMYK"}[im.mode]
    raw = src.read_bytes()
    w, h = im.size
    obj = (f"<< /Type /XObject /Subtype /Image /Width {w} /Height {h} /ColorSpace {cs} "
           f"/BitsPerComponent 8 /Filter /DCTDecode /Length {len(raw)} >>\nstream\n").encode() + raw + b"\nendstream"
    dst.write_bytes(_pdf(w, h, obj))


def img_to_pdf(src, dst):
    from PIL import Image
    im = Image.open(src)
    im.load()
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
        if im.getextrema()[3][0] < 255:        # 투명한 부분이 있으면 흰 바탕에 올린다(보이는 그대로)
            bg = Image.new("RGB", im.size, "white")
            bg.paste(im, mask=im.split()[3])
            im = bg
        else:
            im = im.convert("RGB")
    elif im.mode not in ("RGB", "L"):
        im = im.convert("RGB")
    cs = "/DeviceRGB" if im.mode == "RGB" else "/DeviceGray"
    data = zlib.compress(im.tobytes(), 9)
    w, h = im.size
    obj = (f"<< /Type /XObject /Subtype /Image /Width {w} /Height {h} /ColorSpace {cs} "
           f"/BitsPerComponent 8 /Filter /FlateDecode /Length {len(data)} >>\nstream\n").encode() + data + b"\nendstream"
    dst.write_bytes(_pdf(w, h, obj))
    return None, im


def html_to_pdf(src, dst):
    r = subprocess.run([CHROME, "--headless", "--disable-gpu", "--no-pdf-header-footer",
                        "--run-all-compositor-stages-before-draw", "--virtual-time-budget=8000",
                        f"--print-to-pdf={dst}", src.resolve().as_uri()],
                       capture_output=True, timeout=120)
    if not dst.exists() or dst.stat().st_size == 0:
        return "Chrome 인쇄 실패"


def hwp_to_pdf(src, dst):
    td = Path(tempfile.mkdtemp())
    shutil.copy2(src, td / "in.hwp")
    r = subprocess.run([str(PYHWP / "hwp5html"), "--output", str(td / "h"), str(td / "in.hwp")],
                       capture_output=True, text=True, timeout=600)
    page = td / "h" / "index.xhtml"
    if not page.exists():
        return "pyhwp 변환 실패: " + (r.stderr.strip().splitlines() or ["?"])[-1][:120]
    err = html_to_pdf(page, dst)
    return err or ("hwp_html", page)


def hwpx_to_pdf(src, dst):
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td) / "in.hwpx"
        shutil.copy2(src, tmp)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", td, str(tmp)],
                       capture_output=True, timeout=600)
        out = Path(td) / "in.pdf"
        if not out.exists():
            return "LibreOffice 변환 실패"
        shutil.move(out, dst)


# ---------- 검사 ----------
def _norm(v):
    if v is None:
        return ""
    if isinstance(v, float):
        if v != v:
            return ""
        return repr(round(v, 9)).rstrip("0").rstrip(".") if v != int(v) else str(int(v))
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()[:19]
    s = str(v).strip()
    return s


def _sheets(p):
    """시트별 칸 값 {(시트순서, 행, 열): 값} — 빈 칸 제외."""
    out = {}
    if p.suffix.lower() == ".xls":
        import xlrd
        try:
            wb = xlrd.open_workbook(str(p))
        except xlrd.biffh.XLRDError as e:
            if "encrypted" not in str(e):
                raise
            # 엑셀 기본 암호(VelvetSweatshop, 비밀번호 없이 열리는 '읽기 보호')만 푼다
            import msoffcrypto
            of = msoffcrypto.OfficeFile(open(p, "rb"))
            of.load_key(password="VelvetSweatshop")
            dec = Path(tempfile.mkdtemp()) / "dec.xls"
            with open(dec, "wb") as fh:
                of.decrypt(fh)
            wb = xlrd.open_workbook(str(dec))
        for si, sh in enumerate(wb.sheets()):
            for r in range(sh.nrows):
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    v = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        v = xlrd.xldate_as_datetime(v, wb.datemode)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        v = bool(v)
                    n = _norm(v)
                    if n:
                        out[(si, r, c)] = n
    else:
        import openpyxl
        wb = openpyxl.load_workbook(str(p), data_only=True)
        for si, ws in enumerate(wb.worksheets):
            for row in ws.iter_rows():
                for cell in row:
                    n = _norm(cell.value)
                    if n:
                        out[(si, cell.row - 1, cell.column - 1)] = n
    return out


def check_xls(src, dst):
    a, b = _sheets(src), _sheets(dst)
    if not a:
        return False, "원본에서 값을 못 읽음"
    diff = [k for k in set(a) | set(b) if a.get(k) != b.get(k)]
    if diff:
        k = sorted(diff)[0]
        return False, f"칸 {len(diff)}개 다름(예: 시트{k[0]+1} {k[1]+1}행{k[2]+1}열 {a.get(k)!r}→{b.get(k)!r})"
    return True, f"칸 값 {len(a)}개 전부 같음"


def check_json(src, dst):
    import openpyxl
    data = json.loads(src.read_text(encoding="utf-8"))
    rows = data if isinstance(data, list) else next((v for v in data.values() if isinstance(v, list)), [data])

    def leaves(x):
        if isinstance(x, dict):
            for v in x.values():
                yield from leaves(v)
        elif isinstance(x, list):
            for v in x:
                yield from leaves(v)
        else:
            yield _norm(x)
    a = Counter(v for v in leaves(rows) if v)
    ws = openpyxl.load_workbook(str(dst)).active
    b = Counter(_norm(c.value) for row in ws.iter_rows(min_row=2) for c in row if _norm(c.value))
    if a != b:
        return False, f"값 개수 다름(원본 {sum(a.values())}, 변환 {sum(b.values())})"
    return True, f"값 {sum(a.values())}개 전부 같음"


def _extract(pdf):
    td = tempfile.mkdtemp()
    subprocess.run(["pdfimages", "-all", str(pdf), f"{td}/x"], capture_output=True)
    return sorted(Path(td).iterdir())


def check_jpg(src, dst):
    ex = _extract(dst)
    ok = len(ex) == 1 and ex[0].read_bytes() == src.read_bytes()
    return ok, "JPEG 바이트 그대로 들어감(재압축 없음)" if ok else "꺼낸 이미지가 원본과 다름"


def check_img(src, dst, im):
    from PIL import Image
    ex = _extract(dst)
    if len(ex) != 1:
        return False, "PDF 안 이미지 수가 1이 아님"
    got = Image.open(ex[0]).convert(im.mode)
    ok = got.size == im.size and got.tobytes() == im.tobytes()
    return ok, f"픽셀 {im.size[0]}×{im.size[1]} 전부 같음" if ok else "픽셀이 원본과 다름"


def _toks(s):
    s = re.sub(r"\s+", " ", s)
    nums = re.findall(r"\d+(?:[.,]\d+)*", s)
    words = re.findall(r"[가-힣]{2,}", s)
    return Counter(nums), Counter(words)


def _lost(text, pdf):
    (an, aw), (bn, bw) = _toks(text), _toks(pdf)
    joined = re.sub(r"\s+", "", pdf)
    ln = [k for k in an if bn[k] < an[k] and k not in joined]
    lw = [k for k in aw if bw[k] < aw[k] and k not in joined]
    return an, aw, ln, lw


def check_hwp(src, dst, page):
    ok, msg = check_html(page, dst)
    txt = subprocess.run([str(PYHWP / "hwp5txt"), str(src)], capture_output=True, text=True, timeout=600).stdout
    pdf = subprocess.run(["pdftotext", str(dst), "-"], capture_output=True, text=True).stdout
    an, aw, ln, lw = _lost(txt.replace("<그림>", ""), pdf)
    ok2 = not ln and not lw
    msg = f"HWP→HTML→PDF: {msg}; hwp5txt 본문 숫자 {len(an)-len(ln)}/{len(an)}·글자 {len(aw)-len(lw)}/{len(aw)}종 PDF에 있음"
    if not ok2:
        msg += f" — 빠짐 {ln[:5]} {lw[:5]}"
    return ok and ok2, msg + " (한컴 화면과 줄바꿈·글꼴은 다를 수 있음)"


def check_hwpx(src, dst):
    import zipfile
    z = zipfile.ZipFile(src)
    xml = " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist()
                   if re.match(r"Contents/section\d+\.xml$", n))
    # 그림에 붙은 숨은 설명글(shapeComment: '그림입니다. 원본 그림의 이름…')은 화면에 안 보이므로 뺀다
    xml = re.sub(r"<hp:shapeComment>.*?</hp:shapeComment>", " ", xml, flags=re.S)
    text = " ".join(re.findall(r"<hp:t(?:\s[^>]*)?>(.*?)</hp:t>", xml, re.S))
    text = re.sub(r"<[^>]+>", " ", text)
    pdf = subprocess.run(["pdftotext", str(dst), "-"], capture_output=True, text=True).stdout
    an, aw, ln, lw = _lost(text, pdf)
    if not an and not aw:
        return False, "HWPX 본문 XML에서 글자를 못 찾음"
    ok = not ln and not lw
    msg = f"HWPX 본문 XML 숫자 {len(an)-len(ln)}/{len(an)}·글자 {len(aw)-len(lw)}/{len(aw)}종 PDF에 있음"
    return ok, msg + ("" if ok else f" — 빠짐 {ln[:5]} {lw[:5]}")


def _pdf_texts(dst):
    """pdftotext 세 가지 읽기 순서(기본·raw·layout). 표 칸 안 줄바꿈으로 단어가 쪼개져도 찾을 수 있게 공백을 뺀 것도 쓴다."""
    outs = [subprocess.run(["pdftotext", *m, str(dst), "-"], capture_output=True, text=True).stdout
            for m in ([], ["-raw"], ["-layout"])]
    return outs[2], [re.sub(r"\s+", "", o) for o in outs]


def check_html(src, dst):
    from bs4 import BeautifulSoup
    raw = src.read_bytes()
    m = re.search(rb'charset=["\']?([\w-]+)', raw[:3000], re.I)
    enc = m.group(1).decode() if m else "utf-8"
    try:
        html = raw.decode(enc, errors="ignore")
    except LookupError:
        html = raw.decode("utf-8", errors="ignore")
    soup = BeautifulSoup(html, "html.parser")
    # 인쇄에 나오지 않는 것: 스크립트·스타일, 숨김 요소, 선택 목록(연도 고르기 등)·입력칸
    for t in soup(["script", "style", "noscript", "template", "head", "select", "option", "input", "textarea", "button"]):
        t.decompose()
    for t in soup.select('[style*="display:none"],[style*="display: none"],[hidden],[aria-hidden="true"]'):
        t.decompose()
    text = soup.get_text(" ")
    layout, joined = _pdf_texts(dst)
    (an, aw), (bn, bw) = _toks(text), _toks(layout)
    lost_n = [k for k in an if bn[k] < an[k] and not any(k in j for j in joined)]
    lost_w = [k for k in aw if bw[k] < aw[k] and not any(k in j for j in joined)]
    tn, tw = len(an), len(aw)
    if tn + tw == 0:
        return False, "HTML에 보이는 글자가 없음"
    rw = 1 - len(lost_w) / tw if tw else 1
    msg = f"숫자 {tn - len(lost_n)}/{tn}종, 글자 {tw - len(lost_w)}/{tw}종 PDF에 있음"
    ok = not lost_n and not lost_w          # 하나라도 빠지면 실패(98% 기준은 실제 누락을 통과시켰다)
    if lost_n or lost_w:
        msg += f" — 빠진 숫자 예: {lost_n[:5]} 빠진 글자 예: {lost_w[:5]}"
    return ok, msg


# ---------- 실행 ----------
def rewrite_manifest(path, fn):
    rows = list(csv.DictReader(open(path, encoding="utf-8-sig")))
    rows = [fn(r) for r in rows]
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=R.FIELDS)
        w.writeheader()
        w.writerows(rows)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", default="")
    ap.add_argument("--dry", action="store_true")
    ap.add_argument("--years", default=",".join(R.YEARS))
    ap.add_argument("--skip", default="", help="건드리지 않을 manifest 이름(작업 중인 작업자)")
    a = ap.parse_args()
    only = set(a.only.split(",")) - {""}
    skip = set(a.skip.split(",")) - {""}
    report = []
    for year in a.years.split(","):
        mdir = R.RAW / year / "_manifests"
        for mp in sorted(mdir.glob("*.csv")):
            if mp.stem in skip or mp.stem == WHO:
                continue
            for r in list(csv.DictReader(open(mp, encoding="utf-8-sig"))):
                if r["상태"] != "확보" or not r["파일"] or "/_원본형식/" in r["파일"]:
                    continue
                src = R.RAW / r["파일"]
                ext = src.suffix.lower()
                if ext in KEEP:
                    continue
                kind = KIND.get(ext, "기타")
                if only and kind not in only:
                    continue
                if kind == "기타":
                    report.append((year, r["대학명"], r["파일"], "건너뜀", f"형식 {ext}"))
                    continue
                if not src.exists():
                    report.append((year, r["대학명"], r["파일"], "실패", "원본 파일 없음"))
                    continue
                new_ext = ".xlsx" if kind in ("xls", "json") else ".pdf"
                dst = src.with_name(f"{src.stem}_변환{new_ext}")
                if a.dry:
                    report.append((year, r["대학명"], r["파일"], "예정", dst.name))
                    continue
                im = None
                try:
                    if kind == "xls":
                        err = xls_to_xlsx(src, dst)
                    elif kind == "json":
                        err = json_to_xlsx(src, dst)
                    elif kind == "jpg":
                        err = jpg_to_pdf(src, dst)
                    elif kind == "img":
                        err, im = img_to_pdf(src, dst)
                    elif kind == "hwp":
                        err = hwp_to_pdf(src, dst)
                        if isinstance(err, tuple):
                            im, err = err[1], None
                    elif kind == "hwpx":
                        err = hwpx_to_pdf(src, dst)
                    else:
                        err = html_to_pdf(src, dst)
                    if err:
                        ok, msg = False, err
                    elif kind == "xls":
                        ok, msg = check_xls(src, dst)
                    elif kind == "json":
                        ok, msg = check_json(src, dst)
                    elif kind == "jpg":
                        ok, msg = check_jpg(src, dst)
                    elif kind == "img":
                        ok, msg = check_img(src, dst, im)
                    elif kind == "hwp":
                        ok, msg = check_hwp(src, dst, im)
                    elif kind == "hwpx":
                        ok, msg = check_hwpx(src, dst)
                    else:
                        ok, msg = check_html(src, dst)
                except Exception as e:
                    ok, msg = False, f"오류: {type(e).__name__}: {e}"[:200]
                if not ok:
                    if dst.exists():
                        dst.unlink()
                    report.append((year, r["대학명"], r["파일"], "실패", msg))
                    continue
                keep_dir = src.parent / "_원본형식"
                keep_dir.mkdir(exist_ok=True)
                moved = keep_dir / src.name
                shutil.move(src, moved)
                old_rel, new_rel = r["파일"], str(moved.relative_to(R.RAW))
                conv_rel = str(dst.relative_to(R.RAW))

                def fix(x, old_rel=old_rel, new_rel=new_rel, conv_rel=conv_rel):
                    if x["파일"] == old_rel:
                        x["파일"] = new_rel
                        x["비고"] = (x["비고"] + f" / 형식 통일(2026-09-18): 변환본 {Path(conv_rel).name}").strip(" /")
                    return x
                rewrite_manifest(mp, fix)
                R.append(year, WHO, {
                    "univId": r["univId"], "대학명": r["대학명"], "연도": year, "문서종류": r["문서종류"],
                    "상태": "확보", "형식": new_ext[1:].upper(), "원문여부": r["원문여부"], "파일": conv_rel,
                    "출처URL": r["출처URL"], "게시일": r["게시일"], "받은날짜": dt.date.today().isoformat(),
                    "sha256": R.sha256(dst), "쪽수": R.pages(dst), "작업자": WHO,
                    "비고": f"형식 변환본(원본 {new_rel}, 작업자 {r['작업자']}). 검사: {msg}",
                })
                report.append((year, r["대학명"], conv_rel, "변환", msg))
    out = R.RAW / "_work" / f"형식변환_보고_{dt.date.today().isoformat()}.tsv"
    with open(out, "a", encoding="utf-8") as f:
        for x in report:
            f.write("\t".join(map(str, x)) + "\n")
    c = Counter(x[3] for x in report)
    print(dict(c), "→", out)
    for x in report:
        if x[3] == "실패":
            print("실패", *x[:3], x[4])


if __name__ == "__main__":
    sys.exit(main())
